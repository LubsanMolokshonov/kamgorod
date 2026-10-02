#!/usr/bin/env python3
"""Join every source row to separate production baseline and local acceptance evidence."""
import argparse,csv,json,re
from collections import Counter,defaultdict
from pathlib import Path
from urllib.parse import urlsplit
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment

def normalized(u):
    return 'https://fgos.pro'+(urlsplit(u).path or '/')
def metadata_key(value):
    return re.sub(r'\s+',' ',str(value)).strip().casefold().replace('ё','е')
def summarize(folder):
    crawl=json.loads((folder/'crawl.json').read_text())
    urls=list(csv.DictReader((folder/'url-results.csv').open(encoding='utf-8-sig')))
    sitemap=[r for r in urls if r['sitemap']=='True']
    stats={'crawled':len(crawl),'network_errors':sum(r['http']==0 for r in crawl.values()),'sitemap':len(sitemap),
      'sitemap_not_200':sum(r['http']!='200' for r in sitemap),'sitemap_noindex':sum('noindex' in r['robots'] for r in sitemap),
      'sitemap_noncanonical':sum(len(json.loads(r['canonical']))!=1 or normalized(json.loads(r['canonical'])[0])!=r['url'] for r in sitemap),
      'missing_alt':sum(r['missing_alt'] for r in crawl.values()),
      'broken_internal_pairs':dict(Counter(r['http'] for r in csv.DictReader((folder/'broken-links.csv').open(encoding='utf-8-sig'))))}
    return crawl,stats

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--directory',required=True);a=ap.parse_args();out=Path(a.directory)
    before,bs=summarize(out/'production-before');after,afterstats=summarize(out/'local-after')
    sources=json.loads((out/'local-after/source-registry.json').read_text());content=json.loads((out/'content-decisions.json').read_text())
    editorial=defaultdict(list)
    for row in content:editorial[(row[0],row[1])].append(row)
    indices={f:defaultdict(set) for f in ['title','h1','description']}
    for u,r in after.items():
        if not r['indexable'] or '/page/' in u or len(r['canonical'])!=1 or normalized(r['canonical'][0])!=u:continue
        for f in indices:
            value=' '.join(r[f]) if f=='h1' else r[f]
            if value:indices[f][metadata_key(value)].add(u)
    wb=Workbook();ws=wb.active;ws.title='Все исходные строки'
    ws.append(['Задача','Лист','Строка','Основной URL','Production HTTP до','Локально HTTP','Локально canonical','Индексируемость локально','Входящие HTML локально','Глубина от раздела','Решение','Изменение / основание','Конечный URL','Приёмка','Исходная строка'])
    totals=Counter();orphan_results=Counter()
    for row in sources:
        sheet=row['sheet'];v=row['values'];i={'Сироты':0,'Новые страницы':0,'Дубли метаданных':3,'Падение запросов':4,'Каннибализация':1}[sheet]
        url=normalized(v[i]);b=before.get(url,{});r=after.get(url,{})
        status='нужны подтверждённые данные';reason='Требуется предметная редакционная проверка';result='Не принято';target=url
        matched=editorial.get((sheet,row['row']),[])
        if matched:status=matched[0][3];reason=matched[0][4]
        if sheet=='Сироты':
            if r.get('http')==200 and r.get('incoming',0)>0 and r.get('depth_from_section') is not None and r['depth_from_section']<=3:
                status='уже выполнено' if b.get('incoming',0)>0 and b.get('depth_from_section') is not None and b['depth_from_section']<=3 else 'исправить';result='Локально PASS';reason='HTML-ссылка подтверждена обходом без JavaScript; глубина <=3'
            elif r.get('http') in [301,302]:reason='Проверить исходный и конечный материал; перенаправление не считается карточкой';target=r.get('location','')
            elif r.get('http')==404:reason='В локальной БД/маршрутах URL отсутствует; сопоставить production и актуальное предложение'
            else:reason='Недостаточно входящих ссылок или глубина >3';status='исправить'
            orphan_results[result]+=1
        elif sheet=='Новые страницы':
            ok=r.get('http')==200 and r.get('title')==v[4] and r.get('h1')==[v[5]] and r.get('indexable') and len(r.get('canonical',[]))==1 and normalized(r['canonical'][0])==url and r.get('incoming',0)>0
            status='исправить';reason='Реализованы точные Title/H1 и обязательные блоки';result='Локально PASS' if ok else 'Требуется проверка'
        elif sheet=='Дубли метаданных' and r.get('http')==200:
            f={'Title':'title','H1':'h1','Description':'description'}.get(v[1]);value=' '.join(r.get('h1',[])) if f=='h1' else r.get(f,'')
            unique=f and len(indices[f][metadata_key(value)])==1 and bool(value)
            valid=unique and (f!='description' or 120<=len(value)<=160)
            if valid and not v[4]:status='исправить';result='Локально PASS';reason='Уникальное фактическое значение проверено среди индексируемых canonical-страниц; название записи сохранено'
            elif not v[4]:reason+='; итоговая уникальность/длина не подтверждена'
        if r.get('http')==301:
            target=r.get('location','');dest=after.get(normalized(target),{})
            if dest.get('http')==200:result='Локально 301 → 200; проверить решение по содержанию'
        if r.get('http')==404 and b.get('http')==200:reason+='; production 200, локальная БД не содержит ту же запись'
        totals[result]+=1
        ws.append([row['task'],sheet,row['row'],url,b.get('http'),r.get('http'),'\n'.join(r.get('canonical',[])),r.get('indexable'),r.get('incoming'),r.get('depth_from_section'),status,reason,target,result,json.dumps(v,ensure_ascii=False)])
    groups=defaultdict(list)
    for source in sources:
        if source['sheet']=='Падение запросов':groups[normalized(source['values'][4])].append(source)
    group_sheet=wb.create_sheet('Группы запросов');group_sheet.append(['Цель из ТЗ','Основной запрос','Дополнительные запросы','Строки','Примечание'])
    for url,items in sorted(groups.items()):
        main=[r['values'][0] for r in items if 'Основной' in str(r['values'][8])]
        extra=[r['values'][0] for r in items if 'Основной' not in str(r['values'][8])]
        group_sheet.append([url,'\n'.join(main),'\n'.join(extra),', '.join(str(r['row']) for r in items),'Точные заголовки пяти новых посадочных приоритетны. Спорные назначения и обещания отмечены на первом листе.'])
    queue=wb.create_sheet('Редакторская очередь');queue.append(['Источник','Строка','URL','Решение','Основание'])
    for row in content:queue.append(row)
    changes=wb.create_sheet('Сравнение сред');changes.append(['Метрика','Production до','Локальная проверка'])
    for key in bs:changes.append([key,json.dumps(bs[key],ensure_ascii=False),json.dumps(afterstats[key],ensure_ascii=False)])
    changes.append(['Ограничение','Снимок production через SSH, без изменения сайта','Старая локальная БД: сравнение не доказывает исправление production'])
    links=wb.create_sheet('Проблемные ссылки локально');links.append(['Источник','Ссылка','HTTP','Location','Решение'])
    for row in csv.DictReader((out/'local-after/broken-links.csv').open(encoding='utf-8-sig')):links.append([row[k] for k in ['source','target','http','location','action']])
    pending=wb.create_sheet('Неполные данные');pending.append(['Объект','Что необходимо'])
    pending.append(['URL и правила; Проблемные ссылки','Листы отсутствуют. Свежий обход — новый реестр, он не восстанавливает оригинальные 80/113/9/863 строки.'])
    pending.append(['FGOS-TZ-11','Подтверждённые автор/составитель/редактор, профиль, квалификация, дата содержательного изменения и источники для материалов без этих сведений. Не назначены автоматически.'])
    pending.append(['Полный цикл оплаты','Проверка реальной оплаты и выдачи документов не выполнялась; нужны изолированные тестовые реквизиты и сценарии до выпуска.'])
    for sh in wb:
        sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
        for cell in sh[1]:cell.font=Font(name='Arial',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='243B53')
        for col in sh.columns:
            letter=col[0].column_letter;sh.column_dimensions[letter].width=55 if letter not in ['A','B','C'] else 24
        for cells in sh.iter_rows(min_row=2):
            for cell in cells:cell.font=Font(name='Arial',size=10);cell.alignment=Alignment(vertical='top',wrap_text=True)
    wb.save(out/'Реестр решений.xlsx')
    report={'source_rows':len(sources),'source_sheets':dict(Counter(r['sheet'] for r in sources)),'production_before':bs,'local_after':afterstats,'row_acceptance':dict(totals),'orphan_acceptance':dict(orphan_results)}
    (out/'acceptance-summary.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False))
if __name__=='__main__':main()
