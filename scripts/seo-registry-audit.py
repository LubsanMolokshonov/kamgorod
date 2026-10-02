#!/usr/bin/env python3
"""Read-only GET audit. Registry provenance and crawl graph are kept separately."""
import argparse, csv, json, re, time
from collections import defaultdict, deque
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import urljoin, urlsplit
import xml.etree.ElementTree as ET
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

SKIP = re.compile(r'^/(admin|api|ajax|uploads|assets|login|logout|register|cart|checkout|kabinet|lichnyy-kabinet|opublikovat|sertifikat-publikacii|ai-chat|download)(/|\.|$)')

def registry(path):
    records, urls = [], set()
    book = load_workbook(path, read_only=True, data_only=True)
    tasks = {'Сироты':'FGOS-TZ-14','Каннибализация':'FGOS-TZ-15','Новые страницы':'FGOS-TZ-01,12','Дубли метаданных':'FGOS-TZ-06,13','Падение запросов':'Запросы'}
    for sheet in book:
        for number, row in enumerate(sheet.iter_rows(min_row=6, values_only=True), 6):
            if not any(x is not None for x in row): continue
            links = sorted(set(u.rstrip('.,;') for cell in row if isinstance(cell,str) for u in re.findall(r'https://fgos\.pro/[^\s<>«»\";]*', cell)))
            urls.update(links)
            records.append({'task':tasks.get(sheet.title,sheet.title),'sheet':sheet.title,'row':number,'values':list(row),'urls':links})
    return records, urls

def crawl(base, seeds, workers, limit, previous=None):
    results = dict(previous or {}); pending = set(seeds); pending.add('https://fgos.pro/')
    def fetch(url):
        parts = urlsplit(url)
        path = (parts.path or '/') + ('?' + parts.query if parts.query else '')
        result = {'url':url,'http':0,'location':'','canonical':[],'robots':[],'title':'','h1':[],'description':'','missing_alt':0,'links':[],'error':''}
        try:
            r = requests.get(base.rstrip('/')+path, timeout=(8,35), allow_redirects=False, headers={'User-Agent':'FGOS-Registry-QA/2.0'})
            result['http']=r.status_code
            result['location']=r.headers.get('Location','')
            result['robots']=[r.headers.get('X-Robots-Tag','')]
            if r.status_code==200 and 'text/html' in r.headers.get('Content-Type',''):
                s=BeautifulSoup(r.content,'html.parser')
                result.update(title=s.title.get_text(' ',strip=True) if s.title else '',h1=[h.get_text(' ',strip=True) for h in s.find_all('h1')],
                    canonical=[urljoin(url,x.get('href','')) for x in s.select('link[rel=canonical]')],
                    description=next((x.get('content','') for x in s.select('meta[name=description]')),''),
                    missing_alt=len(s.select('img:not([alt])')))
                result['robots'] += [x.get('content','') for x in s.select('meta[name=robots]')]
                for a in s.select('a[href]'):
                    target=urlsplit(urljoin(url,a['href']))
                    if target.hostname in ['fgos.pro','localhost','127.0.0.1'] and not target.query and not SKIP.search(target.path) and not re.search(r'\.(pdf|jpg|png|zip|docx?|svg|mp4)$',target.path,re.I):
                        result['links'].append('https://fgos.pro'+(target.path or '/'))
                result['links']=sorted(set(result['links']))
        except requests.RequestException as e: result['error']=str(e)
        return result
    with ThreadPoolExecutor(max_workers=workers) as pool:
        while pending and len(results)<limit:
            pending.difference_update(results)
            if not pending: break
            batch=sorted(pending)[:min(60,limit-len(results))];pending.difference_update(batch)
            for item in pool.map(fetch,batch):
                results[item['url']]=item
                discover = list(item['links'])
                if item['location']:
                    destination=urlsplit(urljoin(item['url'],item['location']))
                    if destination.hostname in ['fgos.pro','localhost','127.0.0.1'] and not SKIP.search(destination.path) and not destination.query:
                        discover.append('https://fgos.pro'+destination.path)
                for link in discover:
                    if link not in results: pending.add(link)
            print('Обход:',len(results),'в очереди:',len(pending),flush=True)
    incoming=defaultdict(set)
    for url,row in results.items():
        for target in row['links']:
            if target!=url: incoming[target].add(url)
    depth={u:0 for u in ['https://fgos.pro/','https://fgos.pro/konkursy/','https://fgos.pro/olimpiady/','https://fgos.pro/kursy/','https://fgos.pro/publikacii/','https://fgos.pro/materialy/','https://fgos.pro/vebinary/'] if u in results}
    queue=deque(depth)
    while queue:
        u=queue.popleft()
        for v in results.get(u,{}).get('links',[]):
            if v not in depth: depth[v]=depth[u]+1;queue.append(v)
    for url,row in results.items():
        chain=[];seen={url};cursor=url;loop=False
        while results.get(cursor,{}).get('http') in [301,302,303,307,308]:
            location=results[cursor].get('location','')
            if not location:break
            destination=urlsplit(urljoin(cursor,location))
            target=('https://fgos.pro'+destination.path) if destination.hostname in ['fgos.pro','localhost','127.0.0.1'] and not destination.query else urljoin(cursor,location)
            chain.append(target)
            if target in seen:loop=True;break
            seen.add(target);cursor=target
        row['redirect_chain']=chain;row['redirect_loop']=loop;row['final_url']=cursor;row['final_http']=results.get(cursor,{}).get('http')
        row['incoming']=len(incoming[url]);row['depth_from_section']=depth.get(url)
        row['indexable']=row['http']==200 and not any('noindex' in x.lower() for x in row['robots'])
    return results, bool(pending)

def main():
    p=argparse.ArgumentParser();p.add_argument('--registry',required=True);p.add_argument('--base',required=True);p.add_argument('--output',required=True);p.add_argument('--workers',type=int,default=3);p.add_argument('--limit',type=int,default=10000);p.add_argument('--refresh',nargs='*',help='Повторить указанные пути поверх существующего обхода');a=p.parse_args()
    out=Path(a.output);out.mkdir(parents=True,exist_ok=True)
    rows,urls=registry(a.registry)
    (out/'source-registry.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2))
    sitemap=requests.get(a.base.rstrip('/')+'/sitemap.xml',timeout=90);sitemap.raise_for_status()
    locs={'https://fgos.pro'+urlsplit(x.text).path for x in ET.fromstring(sitemap.content).iter('{http://www.sitemaps.org/schemas/sitemap/0.9}loc')}
    urls.update(locs)
    previous=None
    if a.refresh is not None:
        previous=json.loads((out/'crawl.json').read_text())
        for path in a.refresh:
            url='https://fgos.pro'+urlsplit(path).path
            previous.pop(url,None);urls.add(url)
    results,truncated=crawl(a.base,urls,a.workers,a.limit,previous)
    (out/'crawl.json').write_text(json.dumps(results,ensure_ascii=False,indent=2))
    with (out/'url-results.csv').open('w',newline='',encoding='utf-8-sig') as f:
        fields=['url','http','location','canonical','robots','title','h1','description','missing_alt','incoming','depth_from_section','indexable','sitemap','error','redirect_chain','redirect_loop','final_url','final_http'];w=csv.DictWriter(f,fieldnames=fields);w.writeheader()
        for url,r in sorted(results.items()):w.writerow({k:(json.dumps(r[k],ensure_ascii=False) if isinstance(r.get(k),list) else r.get(k,'')) if k!='sitemap' else url in locs for k in fields})
    with (out/'broken-links.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f);w.writerow(['source','target','http','location','action'])
        for url,r in results.items():
            for link in r['links']:
                target=results.get(link,{})
                if target.get('http',0)!=200:w.writerow([url,link,target.get('http'),target.get('location'),'Проверить конечную цель' if target.get('location') else 'Проверить содержимое; не удалять автоматически'])
    with (out/'decisions.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f);w.writerow(['task','sheet','row','urls','status','reason'])
        for row in rows:
            data=[results.get(u,{}) for u in row['urls']]
            status='нужны подтверждённые данные';reason='Требуется проверка содержания и назначения строки'
            if row['sheet']=='Сироты' and data:
                primary=results.get(row['values'][0],{})
                if primary.get('http')==200 and primary.get('incoming',0)>0 and primary.get('depth_from_section',999) is not None and primary.get('depth_from_section',999)<=3:status='уже выполнено';reason='Есть входящая HTML-ссылка, глубина от раздела <=3'
                elif primary.get('http')==200:status='исправить';reason='Не подтверждена доступность за <=3 перехода'
            w.writerow([row['task'],row['sheet'],row['row'],'\n'.join(row['urls']),status,reason])
    summary={'transport':a.base,'production':a.base!='http://localhost:8080','registry_rows':len(rows),'sitemap_urls':len(locs),'crawled':len(results),'truncated':truncated,'errors':sum(r['http']==0 for r in results.values()),'missing_sheets':['URL и правила','Проблемные ссылки']}
    (out/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2));print(json.dumps(summary,ensure_ascii=False))
if __name__=='__main__':main()
