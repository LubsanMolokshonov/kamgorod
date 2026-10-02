#!/usr/bin/env python3
"""Prepare fact-based SEO overrides; never modifies the database."""
import argparse, hashlib, json, re, unicodedata
from collections import defaultdict
from pathlib import Path
from urllib.parse import urlsplit
from bs4 import BeautifulSoup
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def plain(s): return re.sub(r'\s+',' ',BeautifulSoup(s or '', 'html.parser').get_text(' ',strip=True)).strip()
def norm(s): return plain(s).casefold().replace('ё','е')
def digest(row): return hashlib.sha256((str(row.get('title',''))+'\n'+str(row.get('content') or '')).encode()).hexdigest()

def prepare(snapshot, book):
    entities={}; packages={}; decisions=[]
    for table,prefix in [('competitions','konkursy'),('olympiads','olimpiady'),('courses','kursy'),('materials','material'),('publications','publikaciya')]:
        for row in snapshot[table]:
            section='blog' if table=='publications' and row.get('source')=='blog' else prefix
            if row.get('slug'):entities['/'+section+'/'+row['slug']+'/']=(table,row)
    duplicate_rows=list(book['Дубли метаданных'].values)[5:]
    by_path=defaultdict(list)
    for number,row in enumerate(duplicate_rows,6):by_path[urlsplit(row[3]).path].append((number,row))
    titles=defaultdict(list)
    for path in by_path:
        if path in entities:titles[norm(entities[path][1]['title'])].append(path)
    for path,rows in by_path.items():
        if path not in entities:
            for n,row in rows:decisions.append(['Дубли метаданных',n,path,'нужны подтверждённые данные','URL отсутствует в опубликованном снимке'])
            continue
        table,entity=entities[path]; title=plain(entity['title']);heading=title
        if len(titles[norm(title)])>1:
            author=plain(entity.get('public_author') or entity.get('author_name') or '')
            if author:heading += ' — '+author
            elif table in ['competitions','olympiads']:
                audience={'teachers':'для педагогов','students':'для школьников','preschoolers':'для дошкольников'}.get(entity.get('target_audience'),'')
                subject=plain(entity.get('subject'))
                suffix=', '.join(v for v in [subject,audience] if v and norm(v) not in norm(title))
                if suffix:heading+=' — '+suffix
        desc=plain(entity.get('annotation') or entity.get('description') or entity.get('content') or '')
        if len(desc)<120:
            action={'courses':'Изучите программу, объём часов, условия обучения и оформления итогового документа.', 'competitions':'Ознакомьтесь с условиями конкурса, требованиями к работе и порядком оформления диплома.', 'olympiads':'Изучите условия олимпиады, аудиторию и порядок прохождения заданий.', 'publications':'Прочитайте авторский педагогический материал и ознакомьтесь с содержанием публикации.', 'materials':'Ознакомьтесь с педагогическим материалом и условиями его использования.'}[table]
            desc=(title+'. '+action)
        if len(desc)>160:desc=desc[:157].rsplit(' ',1)[0].rstrip(' ,;:')+'…'
        package={'path':path,'table':table,'id':entity['id'],'expected_sha256':digest(entity),'expected':{k:entity[k] for k in ['title','slug','content','description','annotation','user_id','source','status','is_active','public_author'] if k in entity},'meta_title':heading+' | Педагогический портал','seo_h1':heading,'meta_description':desc}
        if len(desc)<120:package.pop('meta_description')
        packages[path]=package
        for number,row in rows:
            reason='Независимые SEO-поля по фактическому названию и содержанию; требуется HTTP-приёмка'
            status='исправить'
            if row[4]:
                target=entities.get(urlsplit(row[4]).path)
                same=target and table==target[0] and entity.get('user_id')==target[1].get('user_id') and norm(entity.get('content') or entity.get('description')) and norm(entity.get('content') or entity.get('description'))==norm(target[1].get('content') or target[1].get('description'))
                status='нужны подтверждённые данные' if same else 'ошибка исходного требования'
                reason='Совпадение автора и содержимого: дополнительно проверить документы перед 301' if same else 'Технический дубль не доказан; исходный URL и авторство сохранены'
            decisions.append(['Дубли метаданных',number,path,status,reason])
    # Не создаём новые дубли и не маскируем их номерами записей.
    for field in ['seo_h1','meta_title','meta_description']:
        groups=defaultdict(list)
        for path,package in packages.items():
            if package.get(field):groups[norm(package[field])].append(path)
        for paths in groups.values():
            if len(paths)>1:
                for path in paths:
                    packages[path].pop(field,None)
                    decisions.append(['Редакторская очередь','',path,'нужны подтверждённые данные','Нужно содержательное различие для '+field+'; шаблонный дубль не применён'])
    for sheet,query_col,target_col in [('Падение запросов',0,4),('Каннибализация',0,1)]:
        for number,row in enumerate(list(book[sheet].values)[5:],6):
            path=urlsplit(row[target_col]).path;query=row[query_col]
            status='нужны подтверждённые данные';reason='Проверить покрытие запроса содержанием и анкорами: '+query
            if path=='/zhurnal/' and any(word in query for word in ['напечатать','опубликовать','разместить']):
                status='исправить';reason='Коммерческое размещение: /publikaciya-dlya-pedagogov/; журнал сохраняет чтение материалов'
            elif 'конкурс' in query and path.startswith('/olimpiady/'):
                status='ошибка исходного требования';reason='Запрос конкурса назначен олимпиаде; сохраняем разные виды мероприятий'
            elif 'сертификат' in query and path.startswith('/olimpiady/'):
                status='ошибка исходного требования';reason='Для олимпиады проверяется фактический итоговый документ, не обещаем сертификат вместо диплома'
            elif path in ['/kursy/','/kursy/povyshenie-kvalifikatsii/'] and 'воспитател' in query:
                status='исправить';reason='Узкий запрос воспитателей направить в /kursy/povyshenie-kvalifikatsii/dou/; общий каталог не переименовывать'
            elif 'переподготов' in query and '/povyshenie-kvalifikatsii/' in path:
                status='ошибка исходного требования';reason='Переподготовка не относится к КПК; требуется соответствующее предложение'
            elif path=='/kursy/nachalnaya-shkola/' and 'повышения квалификации' in query:
                status='исправить';reason='Цель КПК: /kursy/povyshenie-kvalifikatsii/nachalnaya-shkola/; смешанный каталог сохраняется'
            elif re.search(r'202[0-5]|бесплат',query):
                status='нужны подтверждённые данные';reason='Не переносить устаревший год или неподтверждённое бесплатное предложение в заголовок'
            decisions.append([sheet,number,path,status,reason])
    return list(packages.values()),decisions

def main():
    p=argparse.ArgumentParser();p.add_argument('--snapshot',required=True);p.add_argument('--registry',required=True);p.add_argument('--output',required=True);a=p.parse_args()
    out=Path(a.output);out.mkdir(parents=True,exist_ok=True)
    data=json.loads(Path(a.snapshot).read_text());book=load_workbook(a.registry,read_only=True,data_only=True)
    packages,decisions=prepare(data,book)
    # Контроль окончательных значений по всему production-обходу, включая URL вне Excel.
    baseline=out/'production-before/crawl.json'
    if baseline.exists():
        current=json.loads(baseline.read_text());by_path={row['path']:row for row in packages}
        for field,html_field in [('meta_title','title'),('seo_h1','h1'),('meta_description','description')]:
            groups=defaultdict(set)
            for url,row in current.items():
                path=urlsplit(url).path
                if not row['indexable'] or '/page/' in path or len(row['canonical'])!=1 or urlsplit(row['canonical'][0]).path!=path:continue
                value=by_path.get(path,{}).get(field,row[html_field]);value=' '.join(value) if isinstance(value,list) else value
                if value:groups[norm(value)].add(path)
            for paths in groups.values():
                if len(paths)>1:
                    for path in paths:
                        if field in by_path.get(path,{}):
                            by_path[path].pop(field)
                            decisions.append(['Редакторская очередь','',path,'нужны подтверждённые данные','Итоговое '+field+' совпадает с другой индексируемой страницей production; поле не применяется'])
    packages=[row for row in packages if any(field in row for field in ['meta_title','seo_h1','meta_description'])]
    (out/'metadata-package.json').write_text(json.dumps(packages,ensure_ascii=False,indent=2))
    (out/'content-decisions.json').write_text(json.dumps(decisions,ensure_ascii=False,indent=2))
    wb=Workbook();ws=wb.active;ws.title='Решения';ws.append(['Источник','Строка','URL','Статус','Основание']);
    for row in decisions:ws.append(row)
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for row in ws:
        for c in row:c.font=Font(name='Arial',size=10);c.alignment=Alignment(vertical='top',wrap_text=True)
    for c in ws[1]:c.font=Font(name='Arial',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='334155')
    for col,width in [('A',25),('B',9),('C',65),('D',32),('E',90)]:ws.column_dimensions[col].width=width
    wb.save(out/'Реестр решений.xlsx')
    print(json.dumps({'overrides':len(packages),'decisions':len(decisions)},ensure_ascii=False))
if __name__=='__main__':main()
