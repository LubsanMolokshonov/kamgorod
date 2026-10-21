#!/usr/bin/env python3
"""Сверка приложений с публичным снимком. Пакет JSON, без подключения к БД."""
import argparse
import hashlib
import json
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook

DISPUTED = {39, 72, 79, 119, 130, 133, 149, 155}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def prepare(snapshot, courses_path, competitions_path, brief_path):
    courses = {int(r['id']): r for r in snapshot['courses']}
    competitions = {r['slug']: r for r in snapshot['competitions']}
    overrides = {r['path']: r for r in snapshot['seo_page_overrides']}
    book = load_workbook(courses_path, read_only=True, data_only=True)
    package = {'package_id': 'catalog-20261021', 'source_sha256': {
        Path(p).name: digest(p) for p in [courses_path, competitions_path, brief_path]},
        'source_captured_at': snapshot['captured_at'], 'courses': [], 'metadata': [], 'facts': []}
    report = {'courses': [], 'competitions': [], 'missing_inputs': [
        'Полное ТЗ 012', 'Подробное ТЗ 014', 'Семантика конкурсов', '313 исторических сирот']}
    seen_ids, seen_paths = set(), set()
    for row in list(book['Новые названия'].values)[1:]:
        ident = int(row[1]); current = courses.get(ident); path = urlsplit(row[4]).path
        if ident in seen_ids or path in seen_paths:
            raise ValueError('Повтор ID или URL курса в приложении')
        seen_ids.add(ident); seen_paths.add(path)
        if not isinstance(row[7], str) or not row[7].strip() or len(row[7]) > 255 or not isinstance(row[8], str) or not row[8].strip() or len(row[8]) > 500:
            raise ValueError(f'Некорректное название или Title курса {ident}')
        problems = []
        if ident in DISPUTED: problems.append('Требуется решение владельца программы')
        if not current: problems.append('ID отсутствует в снимке')
        elif (path != '/kursy/' + current['slug'] + '/' or current['title'] != row[5]
              or int(current['hours']) != int(row[3])
              or current['program_type'] != {'КПК': 'kpk', 'ПП': 'pp'}[row[2]]):
            problems.append('Изменились исходные URL, название, часы или тип программы')
        report['courses'].append({'id': ident, 'status': 'hold' if problems else 'ready', 'issues': problems})
        if problems: continue
        expected = {k: current[k] for k in ['title', 'slug', 'hours', 'program_type', 'modules_json', 'is_active']}
        package['courses'].append({'id': ident, 'expected': expected, 'title': row[7], 'path': path})
        old = overrides.get(path, {})
        package['metadata'].append({'path': path, 'entity': 'courses', 'id': ident,
            'expected': {k: old.get(k) for k in ['seo_h1', 'meta_title']},
            'values': {'seo_h1': row[7], 'meta_title': row[8]}})
    book = load_workbook(competitions_path, read_only=True, data_only=True)
    urls = list(book['URL и действия'].values)[4:]
    registry = {'sources': package['source_sha256'], 'categories': {}, 'links': []}
    for row in urls:
        if row[1] == 'категория' and 'KEEP_CATEGORY' in (row[14] or ''):
            registry['categories'][urlsplit(row[0]).path] = {'actions': row[14].split('; ')}
    for row in list(book['Метаданные'].values)[4:]:
        if row[14] != 'Изменить': continue
        path = urlsplit(row[0]).path
        entity = competitions.get(path.strip('/').split('/')[-1])
        old = overrides.get(path, {})
        values = {'seo_h1': row[3], 'meta_title': row[5], 'meta_description': row[7]}
        if path in registry['categories']:
            values.update(intro_text=row[7], index_catalog=1)
        package['metadata'].append({'path': path, 'entity': 'competitions' if entity else 'category',
            'id': entity['id'] if entity else None,
            'expected': {k: old.get(k) for k in values}, 'values': values})
    for row in list(book['Перелинковка и 301'].values)[4:]:
        if row[0] == 'HTML ссылка': registry['links'].append({'from': urlsplit(row[1]).path,
            'to': urlsplit(row[2]).path, 'anchor': row[3], 'placement': row[4]})
    # Явные роли из содержательного ТЗ. Остальные карточки требуют предметной проверки.
    confirmed = {
        'umnyy-malysh-konkurs-dlya-doshkolnikov': (['doshkolnikam'], [], []),
        'my-vmeste-konkurs-dlya-detey-s-ovz': (['doshkolnikam', 'shkolnikam'], [], []),
        'govorim-pravilno-tvorcheskiy-konkurs': (['doshkolnikam', 'shkolnikam'], [], ['logopediya']),
        'sovremennyy-urok-fgos': (['pedagogi'], [], []),
        'logopedicheskaya-masterskaya-innovatsionnye-metodiki': (['pedagogi'], [], ['logopediya']),
    }
    for slug, current in competitions.items():
        roles, stages, subjects = confirmed.get(slug, ([], [], []))
        report['competitions'].append({'id': current['id'], 'slug': slug,
            'target_participants': current['target_participants'], 'work_type': current['category'],
            'status': 'role_confirmed' if roles else 'needs_classification',
            'missing': ['Подтвержденные предметы и ступени', 'Сроки', 'Требования к работе',
                'Утвержденное положение и дата версии']})
        if not roles: continue
        package['facts'].append({'id': current['id'],
            'expected': {k: current[k] for k in ['slug', 'target_participants', 'category']},
            'expected_facts': None,
            'values': {'participant_roles': roles, 'education_stages': stages,
                'subject_slugs': subjects, 'work_type': current['category'], 'applicant_role': 'adult',
                'content_json': {}, 'evidence': 'ТЗ_конкурсы.docx: явное назначение карточки'}})
    return package, registry, report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--snapshot', required=True)
    parser.add_argument('--courses', required=True)
    parser.add_argument('--competitions', required=True)
    parser.add_argument('--brief', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    package, registry, report = prepare(json.loads(Path(args.snapshot).read_text()), args.courses, args.competitions, args.brief)
    out = Path(args.output); out.mkdir(parents=True, exist_ok=True)
    for name, data in [('package.json', package), ('registry.json', registry), ('readiness.json', report)]:
        (out / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'courses_ready': len(package['courses']), 'metadata': len(package['metadata']),
        'facts': len(package['facts']), 'categories': len(registry['categories'])}, ensure_ascii=False))


if __name__ == '__main__':
    main()
